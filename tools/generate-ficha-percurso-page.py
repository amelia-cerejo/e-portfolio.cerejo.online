from __future__ import annotations

from html import escape
from pathlib import Path
import re

from docx import Document


SOURCE = Path(
    r"C:\Users\Cerejo Lopes\Downloads\Formador de Formadores\FCE08\Ficha De Percurso Profissional E De FormaÃ§Ã£o - AmÃ©lia Cerejo Lopes.docx"
)
OUTPUT = Path("evidencias/ficha-percurso.html")


def split_lines(text: str) -> list[str]:
    return [line.strip() for line in text.splitlines() if line.strip()]


def section_id(title: str) -> str:
    cleaned = re.sub(r"^\d+\.\s*", "", title).lower()
    replacements = {
        "Ã§": "c",
        "Ã£": "a",
        "Ã¡": "a",
        "Ã ": "a",
        "Ã¢": "a",
        "Ã©": "e",
        "Ãª": "e",
        "Ã­": "i",
        "Ã³": "o",
        "Ãµ": "o",
        "Ãº": "u",
    }
    for src, dest in replacements.items():
        cleaned = cleaned.replace(src, dest)
    cleaned = re.sub(r"[^a-z0-9]+", "-", cleaned).strip("-")
    return cleaned or "secao"


def display_title(title: str) -> str:
    return re.sub(r"^\d+\.\s*", "", title).strip()


def read_docx() -> tuple[str, list[dict], list[list[str]]]:
    doc = Document(SOURCE)
    title = ""
    sections: list[dict] = []
    current: dict | None = None

    for paragraph in doc.paragraphs:
        for text in split_lines(paragraph.text):
            if not title:
                title = text
                continue

            if re.match(r"^\d+\.\s+", text):
                current = {"title": text, "id": section_id(text), "items": []}
                sections.append(current)
                continue

            if current is None:
                current = {"title": "Notas", "id": "notas", "items": []}
                sections.append(current)

            current["items"].append(text)

    table_rows: list[list[str]] = []
    if doc.tables:
        for row in doc.tables[0].rows:
            table_rows.append([cell.text.strip() for cell in row.cells])

    return title, sections, table_rows


def render_items(section: dict) -> str:
    items = section["items"]
    if not items:
        return ""

    html: list[str] = []
    bullet_buffer: list[str] = []

    def flush_bullets() -> None:
        nonlocal bullet_buffer
        if bullet_buffer:
            html.append("<ul>" + "".join(f"<li>{escape(item)}</li>" for item in bullet_buffer) + "</ul>")
            bullet_buffer = []

    for item in items:
        is_label = ":" in item and len(item.split(":", 1)[0]) <= 38 and section["id"] == "identificacao"
        is_bullet = section["id"] in {"experiencia-na-area-da-formacao", "expectativas-e-objectivos-de-desenvolvimento"} and (
            item.endswith(";") or item.endswith(".")
        )

        if is_label:
            flush_bullets()
            label, value = item.split(":", 1)
            html.append(
                f'<div class="fact"><span>{escape(label)}</span><strong>{escape(value.strip())}</strong></div>'
            )
        elif item.endswith(":") and len(item) < 80:
            flush_bullets()
            html.append(f"<h3>{escape(item)}</h3>")
        elif is_bullet and not item.startswith("Na preparaÃ§Ã£o"):
            bullet_buffer.append(item.rstrip(";"))
        else:
            flush_bullets()
            html.append(f"<p>{escape(item)}</p>")

    flush_bullets()
    return "\n".join(html)


def render_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""

    head = rows[0]
    body = rows[1:]
    return f"""
      <div class="table-wrap">
        <table>
          <thead>
            <tr>{''.join(f'<th>{escape(cell)}</th>' for cell in head)}</tr>
          </thead>
          <tbody>
            {''.join('<tr>' + ''.join(f'<td>{escape(cell)}</td>' for cell in row) + '</tr>' for row in body)}
          </tbody>
        </table>
      </div>
    """


def build_html() -> str:
    title, sections, table_rows = read_docx()
    summary = {
        "Nome": "AmÃ©lia Cerejo Lopes",
        "Ãrea": "TIC, FormaÃ§Ã£o, Psicologia e GestÃ£o",
        "SituaÃ§Ã£o": "Formadora / Trabalhadora independente",
        "Documento": "Ficha de percurso profissional e de formaÃ§Ã£o",
    }

    nav = "\n".join(
        f'<a href="#{escape(section["id"])}">{escape(display_title(section["title"]))}</a>'
        for section in sections
    )

    section_html: list[str] = []
    for section in sections:
        body = render_items(section)
        if section["id"] == "competencias-desenvolvidas":
            body += render_table(table_rows)
        section_html.append(
            f"""
            <section id="{escape(section["id"])}">
              <span class="section-number">{escape(section["title"].split(".", 1)[0])}</span>
              <h2>{escape(display_title(section["title"]))}</h2>
              {body}
            </section>
            """
        )

    facts = "\n".join(
        f'<div class="summary-card"><span>{escape(label)}</span><strong>{escape(value)}</strong></div>'
        for label, value in summary.items()
    )

    return f"""<!doctype html>
<html lang="pt-PT">
  <head>
    <meta charset="UTF-8" />
    <meta name="viewport" content="width=device-width, initial-scale=1.0" />
    <title>{escape(title)} - AmÃ©lia Cerejo Lopes</title>
    <style>
      :root {{
        --ink: #172033;
        --muted: #626d80;
        --blue: #002b5b;
        --gold: #bd8a12;
        --paper: #fbf8f4;
        --soft: #edf3f8;
        --line: #e2ded8;
        --white: #fff;
      }}

      * {{ box-sizing: border-box; }}

      body {{
        margin: 0;
        background: var(--paper);
        color: var(--ink);
        font-family: "Segoe UI", Arial, sans-serif;
        line-height: 1.62;
      }}

      header {{
        background: linear-gradient(135deg, #ffffff 0%, #eef4fb 58%, #f8efe0 100%);
        border-bottom: 1px solid var(--line);
      }}

      .hero {{
        max-width: 1120px;
        margin: 0 auto;
        padding: 44px 22px 34px;
      }}

      .hero > a,
      .hero > button {{
        display: none;
      }}

      .back {{
        display: inline-flex;
        margin-bottom: 26px;
        color: var(--blue);
        text-decoration: none;
        font-weight: 700;
      }}

      h1, h2, h3 {{
        font-family: Georgia, "Times New Roman", serif;
        color: #001736;
        letter-spacing: 0;
      }}

      h1 {{
        max-width: 840px;
        margin: 0;
        font-size: clamp(2rem, 4vw, 3.8rem);
        line-height: 1.05;
      }}

      .lead {{
        max-width: 760px;
        margin: 16px 0 0;
        color: var(--muted);
        font-size: 1.05rem;
      }}

      .summary {{
        display: grid;
        grid-template-columns: repeat(4, minmax(0, 1fr));
        gap: 12px;
        margin-top: 28px;
      }}

      .summary-card, .fact {{
        background: rgba(255, 255, 255, 0.82);
        border-left: 3px solid var(--gold);
        padding: 12px 14px;
      }}

      .summary-card span, .fact span {{
        display: block;
        color: var(--muted);
        font-size: 0.74rem;
        font-weight: 800;
        text-transform: uppercase;
        letter-spacing: 0.04em;
      }}

      .layout {{
        max-width: 1120px;
        margin: 0 auto;
        padding: 28px 22px 56px;
        display: grid;
        grid-template-columns: 260px minmax(0, 1fr);
        gap: 24px;
      }}

      nav {{
        position: sticky;
        top: 18px;
        align-self: start;
        background: var(--white);
        border: 1px solid var(--line);
        border-radius: 8px;
        padding: 12px;
      }}

      nav a {{
        display: block;
        padding: 8px 10px;
        border-radius: 6px;
        color: var(--blue);
        text-decoration: none;
        font-weight: 700;
        font-size: 0.91rem;
      }}

      nav a:hover {{
        background: var(--soft);
      }}

      section {{
        position: relative;
        background: var(--white);
        border: 1px solid var(--line);
        border-radius: 8px;
        padding: clamp(22px, 3vw, 34px);
        margin-bottom: 18px;
      }}

      .section-number {{
        display: inline-grid;
        place-items: center;
        width: 34px;
        height: 34px;
        border-radius: 50%;
        background: var(--blue);
        color: #fff;
        font-weight: 800;
        margin-bottom: 12px;
      }}

      h2 {{
        margin: 0 0 14px;
        font-size: clamp(1.35rem, 2vw, 2rem);
      }}

      h3 {{
        margin: 18px 0 8px;
        font-size: 1.08rem;
      }}

      p {{
        margin: 0 0 14px;
      }}

      ul {{
        margin: 8px 0 14px;
        padding-left: 1.2rem;
      }}

      li {{
        margin: 6px 0;
      }}

      .fact {{
        margin: 10px 0;
        border-radius: 6px;
        border-top: 1px solid #f1ede7;
        border-right: 1px solid #f1ede7;
        border-bottom: 1px solid #f1ede7;
      }}

      .table-wrap {{
        overflow-x: auto;
        margin-top: 12px;
      }}

      table {{
        width: 100%;
        border-collapse: collapse;
        background: #fff;
      }}

      th {{
        background: var(--blue);
        color: #fff;
        text-align: left;
        padding: 12px;
        font-size: 0.88rem;
      }}

      td {{
        border: 1px solid var(--line);
        padding: 13px 12px;
        vertical-align: top;
      }}

      td:first-child {{
        width: 28%;
        color: var(--blue);
        font-weight: 800;
      }}

      footer {{
        max-width: 1120px;
        margin: 0 auto;
        padding: 0 22px 34px;
        color: var(--muted);
        font-size: 0.9rem;
      }}

      .floating-nav {{
        position: fixed;
        right: 22px;
        bottom: 22px;
        z-index: 40;
        display: flex;
        flex-direction: column;
        align-items: flex-end;
        gap: 8px;
      }}

      .float-button {{
        width: 46px;
        height: 46px;
        display: grid;
        place-items: center;
        border: 1px solid #cad4e2;
        border-radius: 999px;
        background: var(--white);
        color: var(--blue);
        box-shadow: 0 12px 28px rgba(0, 43, 91, 0.14);
        text-decoration: none;
        padding: 0;
        cursor: pointer;
      }}

      .float-button svg {{
        width: 20px;
        height: 20px;
        stroke: currentColor;
        stroke-width: 2.2;
        stroke-linecap: round;
        stroke-linejoin: round;
        fill: none;
      }}

      .float-button.action-print {{
        background: #dff4e6;
        border-color: #4f9a69;
        color: #155c32;
      }}

      .float-button.action-complete {{
        background: #ffe9a8;
        border-color: #c98a00;
        color: #694400;
      }}

      .float-button:hover {{
        border-color: var(--gold);
        transform: translateY(-1px);
      }}

      .float-button.is-hidden {{
        display: none;
      }}

      .sr-only {{
        position: absolute;
        width: 1px;
        height: 1px;
        padding: 0;
        margin: -1px;
        overflow: hidden;
        clip: rect(0, 0, 0, 0);
        white-space: nowrap;
        border: 0;
      }}

      @media (max-width: 840px) {{
        .summary {{
          grid-template-columns: repeat(2, minmax(0, 1fr));
        }}

        .layout {{
          grid-template-columns: 1fr;
        }}

        nav {{
          position: static;
        }}
      }}

      @media (max-width: 560px) {{
        .summary {{
          grid-template-columns: 1fr;
        }}
      }}

      @media print {{
        .floating-nav, nav {{
          display: none;
        }}

        body {{
          background: #fff;
        }}

        .layout {{
          display: block;
          max-width: none;
          padding: 0;
        }}

        section {{
          border: 0;
          border-bottom: 1px solid #ddd;
          border-radius: 0;
          break-inside: avoid;
        }}
      }}
    </style>
  </head>
  <body>
    <header>
      <div class="hero">
        <h1>{escape(title)}</h1>
        <p class="lead">PÃ¡gina de evidÃªncia complementar com a sÃ­ntese do percurso profissional, formativo e pedagÃ³gico de AmÃ©lia Cerejo Lopes.</p>
        <div class="summary">{facts}</div>
      </div>
    </header>

    <div class="layout">
      <nav aria-label="Ãndice da ficha">
        {nav}
      </nav>
      <main>
        {''.join(section_html)}
      </main>
    </div>

    <footer>
      Fonte: {escape(SOURCE.name)} Â· PÃ¡gina local gerada para o e-PortfÃ³lio.
    </footer>

    <div class="floating-nav" aria-label="Navegacao rapida">
      <a class="float-button" href="../index.html#evidencias" title="Voltar as evidencias" aria-label="Voltar as evidencias">
        <svg viewBox="0 0 24 24" aria-hidden="true"><path d="M19 12H5"/><path d="m12 19-7-7 7-7"/></svg>
        <span class="sr-only">Voltar as evidencias</span>
      </a>
      <button class="float-button action-print" type="button" onclick="window.print()" title="Imprimir esta ficha" aria-label="Imprimir esta ficha">
        <svg viewBox="0 0 24 24" aria-hidden="true"><path d="M7 8V4h10v4"/><path d="M7 17H5a2 2 0 0 1-2-2v-4a2 2 0 0 1 2-2h14a2 2 0 0 1 2 2v4a2 2 0 0 1-2 2h-2"/><path d="M7 14h10v6H7z"/><path d="M17 12h.01"/></svg>
        <span class="sr-only">Imprimir esta ficha</span>
      </button>
      <a class="float-button action-complete" href="../imprimir-completo.html" title="Imprimir e-Portfolio completo" aria-label="Imprimir e-Portfolio completo">
        <svg viewBox="0 0 24 24" aria-hidden="true"><path d="M8 7h10v13H8z"/><path d="M6 17H5a2 2 0 0 1-2-2V4h10a2 2 0 0 1 2 2v1"/><path d="M11 11h4"/><path d="M11 15h4"/></svg>
        <span class="sr-only">Imprimir e-Portfolio completo</span>
      </a>
      <button class="float-button is-hidden" type="button" id="float-top" onclick="scrollToTop()" title="Topo" aria-label="Topo">
        <svg viewBox="0 0 24 24" aria-hidden="true"><path d="M12 19V5"/><path d="m5 12 7-7 7 7"/></svg>
        <span class="sr-only">Topo</span>
      </button>
    </div>

    <script>
      const topButton = document.getElementById("float-top");
      const updateTopButton = () => {{
        if (!topButton) return;
        topButton.classList.toggle("is-hidden", window.scrollY < 180);
      }};
      function scrollToTop() {{
        window.scrollTo({{ top: 0, behavior: "smooth" }});
      }}
      updateTopButton();
      window.addEventListener("scroll", updateTopButton, {{ passive: true }});
    </script>
  </body>
</html>
"""


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(build_html(), encoding="utf-8")
    print(f"PÃ¡gina criada: {OUTPUT}")


if __name__ == "__main__":
    main()

